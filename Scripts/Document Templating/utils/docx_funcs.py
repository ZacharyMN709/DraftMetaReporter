from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import docx.opc.constants
import docx.oxml.ns

# Font Defaults
FONT_NAME = 'Calibri Light'
FONT_SIZE = 16
FONT_COLOR = 0x1F, 0x37, 0x63


def set_cell_margins(cell, **kwargs):
    """
    provided values are in twentieths of a point (1/1440 of an inch).
    read more here: http://officeopenxml.com/WPtableCellMargins.php
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = docx.oxml.shared.OxmlElement('w:tcMar')

    for m in ["top", "start", "bottom", "end"]:
        if m in kwargs:
            node = docx.oxml.shared.OxmlElement("w:{}".format(m))
            node.set(docx.oxml.ns.qn('w:w'), str(kwargs.get(m)))
            node.set(docx.oxml.ns.qn('w:type'), 'dxa')
            tcMar.append(node)

    tcPr.append(tcMar)


def add_image_to_cell(cell, height: float, width: float, file_loc: str):
    set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    paragraph.add_run().add_picture(file_loc, height=Inches(height), width=Inches(width))


def apply_default_font(font):
    apply_font(font)


def apply_font(font, name: str = None, size: int = None, color: tuple[int, int, int] = None):
    font.name = name or FONT_NAME
    font.size = Pt(size or FONT_SIZE)
    font.color.rgb = RGBColor(*(color or FONT_COLOR))


def update_margins(document, top_margin: float, bottom_margin: float, left_margin: float, right_margin: float):
    sections = document.sections
    for section in sections:
        section.top_margin = Cm(top_margin)
        section.bottom_margin = Cm(bottom_margin)
        section.left_margin = Cm(left_margin)
        section.right_margin = Cm(right_margin)


def add_hyperlink(document, run, url: str, is_external: bool = True):
    r = document.part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=is_external)

    r_hyperlink = docx.oxml.shared.OxmlElement('w:hyperlink')
    r_hyperlink.set(docx.oxml.shared.qn('r:id'), r)
    r_hyperlink.set(docx.oxml.shared.qn('w:history'), '1')

    new_run = docx.oxml.shared.OxmlElement('w:r')
    r_properties = docx.oxml.shared.OxmlElement('w:rPr')

    new_run.append(r_properties)
    new_run.text = run.text
    r_hyperlink.append(new_run)

    run._r.getparent().replace(run._r, r_hyperlink)
    apply_default_font(run.font)

